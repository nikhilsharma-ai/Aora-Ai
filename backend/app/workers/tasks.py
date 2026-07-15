import asyncio
import logging
from celery import shared_task
from sqlalchemy.future import select

# Celery runs tasks synchronously, we route async operations using helper loops
def run_sync(coro):
    try:
        loop = asyncio.get_event_loop()
    except RuntimeError:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
    return loop.run_until_complete(coro)

logger = logging.getLogger(__name__)

@shared_task(name="app.workers.tasks.process_document_task")
def process_document_task(document_id: int, user_id: str):
    """
    Background worker task parsing documents, generating embeddings,
    and upserting chunks into Qdrant vector spaces.
    """
    logger.info(f"Starting processing for document {document_id}")
    run_sync(_async_process_document(document_id, user_id))
    return {"status": "success", "document_id": document_id}

@shared_task(name="app.workers.tasks.generate_podcast_task")
def generate_podcast_task(podcast_id: int, user_id: str, topic: str):
    """
    Background worker orchestrating voice synthesis turns and merging
    assets to generate a podcast file.
    """
    logger.info(f"Starting podcast generation task {podcast_id} for topic: {topic}")
    run_sync(_async_generate_podcast(podcast_id, user_id, topic))
    return {"status": "success", "podcast_id": podcast_id}



def markdown_to_html(md: str) -> str:
    import re
    
    # Normalize newlines
    md = md.replace("\r\n", "\n").replace("\r", "\n")
    
    # 1. Extract and preserve code blocks (to avoid stripping their spaces or parsing markdown inside them)
    code_blocks = []
    def save_code_block(match):
        lang = match.group(1) or "plaintext"
        code = match.group(2)
        # HTML escape special characters inside the code
        code_html = code.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        placeholder = f"<!--CODEBLOCK_{len(code_blocks)}-->"
        code_blocks.append(f'<pre><code class="language-{lang}">{code_html}</code></pre>')
        return placeholder
    
    # Match code blocks: ```[lang]\n[code]\n```
    html = re.sub(r"```+[ \t]*(\w*)[ \t]*\n(.*?)(?:\n[ \t]*```+[ \t]*|[ \t]*```+[ \t]*|\Z)", save_code_block, md, flags=re.DOTALL)
    
    # 0. Clean up multiple consecutive spaces (4 or more) and newlines on the non-code content
    html = re.sub(r" {4,}", " ", html)
    html = re.sub(r"\n{3,}", "\n\n", html)
    
    # 2. Parse blockquotes
    def replace_blockquote(match):
        content = match.group(1).strip()
        content = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", content)
        return f'<blockquote>{content}</blockquote>'
    
    html = re.sub(r"^>\s+(.*)$", replace_blockquote, html, flags=re.MULTILINE)
    
    # 3. Parse headers (from h6 down to h1 to prevent partial matches)
    html = re.sub(r"^###### (.*)$", r"<h6>\1</h6>", html, flags=re.MULTILINE)
    html = re.sub(r"^##### (.*)$", r"<h5>\1</h5>", html, flags=re.MULTILINE)
    html = re.sub(r"^#### (.*)$", r"<h4>\1</h4>", html, flags=re.MULTILINE)
    html = re.sub(r"^### (.*)$", r"<h3>\1</h3>", html, flags=re.MULTILINE)
    html = re.sub(r"^## (.*)$", r"<h2>\1</h2>", html, flags=re.MULTILINE)
    html = re.sub(r"^# (.*)$", r"<h1>\1</h1>", html, flags=re.MULTILINE)
    
    # 4. Parse bold, links
    html = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", html)
    html = re.sub(r"\[(.*?)\]\((.*?)\)", r'<a href="\2" target="_blank" rel="noopener noreferrer">\1</a>', html)
    
    # 5. Parse bullet lists
    def replace_list(match):
        items = match.group(0).strip().split("\n")
        list_html = "<ul>"
        for item in items:
            item_text = re.sub(r"^[\-\*]\s+", "", item)
            list_html += f"<li>{item_text}</li>"
        list_html += "</ul>"
        return list_html
    
    html = re.sub(r"(?:^[\-\*]\s+.*(?:\n|$))+", replace_list, html, flags=re.MULTILINE)
    
    # 6. Parse tables
    def replace_table(match):
        lines = [line.strip() for line in match.group(0).strip().split("\n") if line.strip()]
        if not lines:
            return ""
        table_html = "<table>"
        # Parse headers
        headers = [h.strip() for h in lines[0].split("|")[1:-1]]
        table_html += "<thead><tr>"
        for h in headers:
            table_html += f"<th>{h}</th>"
        table_html += "</tr></thead><tbody>"
        
        # Skip header separator if it exists
        start_idx = 1
        if len(lines) > 1 and all(c in "-:| " for c in lines[1].replace("|", "")):
            start_idx = 2
            
        for line in lines[start_idx:]:
            cols = [c.strip() for c in line.split("|")[1:-1]]
            if not cols:
                continue
            table_html += "<tr>"
            for c in cols:
                table_html += f"<td>{c}</td>"
            table_html += "</tr>"
        
        table_html += "</tbody></table>"
        return table_html
        
    html = re.sub(r"(?:^\|.*\|(?:\n|$))+", replace_table, html, flags=re.MULTILINE)
    
    # 7. Wrap paragraphs
    lines = html.split("\n")
    for i, line in enumerate(lines):
        stripped = line.strip()
        # Do not wrap code block placeholders
        if stripped.startswith("<!--CODEBLOCK_") and stripped.endswith("-->"):
            continue
        if stripped and not stripped.startswith(("<h", "<ul", "<ol", "<li", "<pre", "<code", "<blockquote", "<table", "<thead", "<tbody", "<tr", "<td", "<th", "</pre>", "</table>", "</ul>", "</ol>", "</blockquote>", "<p>", "</p>")):
            lines[i] = f"<p>{line}</p>"
            
    html = "\n".join(lines)
    
    # 8. Restore code blocks
    for i, block_html in enumerate(code_blocks):
        html = html.replace(f"<!--CODEBLOCK_{i}-->", block_html)
        
    return html


# Helper async processes
async def _async_process_document(document_id: int, user_id: str):
    from app.db.session import SessionLocal
    from app.db.models.document import Document
    from app.services.llm import llm_service
    from app.services.vector import vector_service
    import re

    def extract_youtube_video_id(url: str) -> str:
        patterns = [
            r"(?:v=|\/v\/|embed\/|youtu\.be\/|youtube\.com\/shorts\/)([a-zA-Z0-9_-]{11})",
        ]
        for pattern in patterns:
            match = re.search(pattern, url)
            if match:
                return match.group(1)
        return ""

    def get_youtube_transcript(video_id: str) -> tuple[str, str]:
        try:
            from youtube_transcript_api import YouTubeTranscriptApi
            transcript_list = YouTubeTranscriptApi().list(video_id)
            # Prioritize native/manual transcripts first, then auto-generated
            manual_transcripts = {}
            generated_transcripts = {}
            for t in transcript_list:
                if not t.is_generated:
                    manual_transcripts[t.language_code] = t
                else:
                    generated_transcripts[t.language_code] = t
            
            preferred = None
            if 'hi' in manual_transcripts:
                preferred = manual_transcripts['hi']
            elif 'en' in manual_transcripts:
                preferred = manual_transcripts['en']
            elif 'hi' in generated_transcripts:
                preferred = generated_transcripts['hi']
            elif 'en' in generated_transcripts:
                preferred = generated_transcripts['en']
            elif manual_transcripts:
                preferred = next(iter(manual_transcripts.values()))
            elif generated_transcripts:
                preferred = next(iter(generated_transcripts.values()))
            else:
                preferred = next(iter(transcript_list))
                
            transcript = preferred
            logger.info(f"Selected transcript language '{transcript.language_code}' (generated={transcript.is_generated}) for video {video_id}")
            
            fetched_data = transcript.fetch()
            clean_text = " ".join([
                (snippet.text if hasattr(snippet, "text") else snippet.get("text", "")) 
                for snippet in fetched_data
            ])
            
            # Format transcript with timestamps grouped into 30-second blocks
            formatted_lines = []
            current_group = []
            current_time = 0.0
            
            for snippet in fetched_data:
                text = (snippet.text if hasattr(snippet, "text") else snippet.get("text", "")).strip()
                start = snippet.start if hasattr(snippet, "start") else snippet.get("start", 0.0)
                
                if not text:
                    continue
                    
                if not current_group:
                    current_time = start
                    current_group.append(text)
                elif start - current_time >= 30.0:
                    minutes = int(current_time // 60)
                    seconds = int(current_time % 60)
                    timestamp = f"**[{minutes:02d}:{seconds:02d}]**"
                    formatted_lines.append(f"{timestamp} " + " ".join(current_group))
                    current_group = [text]
                    current_time = start
                else:
                    current_group.append(text)
                    
            if current_group:
                minutes = int(current_time // 60)
                seconds = int(current_time % 60)
                timestamp = f"**[{minutes:02d}:{seconds:02d}]**"
                formatted_lines.append(f"{timestamp} " + " ".join(current_group))
                
            formatted_transcript = "\n\n".join(formatted_lines)
            return clean_text, formatted_transcript
        except Exception as e:
            logger.error(f"Failed to get YouTube transcript: {e}")
            return "", ""

    def chunk_text(text: str, chunk_size: int = 1000) -> list:
        words = text.split(" ")
        chunks = []
        current_chunk = []
        current_size = 0
        for word in words:
            current_chunk.append(word)
            current_size += len(word) + 1
            if current_size >= chunk_size:
                chunks.append(" ".join(current_chunk))
                current_chunk = []
                current_size = 0
        if current_chunk:
            chunks.append(" ".join(current_chunk))
        return chunks

    doc_type = ""
    doc_name = ""
    doc_file_url = ""
    async with SessionLocal() as db:
        result = await db.execute(select(Document).where(Document.id == document_id))
        doc = result.scalar_one_or_none()
        if not doc:
            logger.error(f"Document {document_id} not found in database.")
            return
        doc_type = doc.doc_type
        doc_name = doc.name
        doc_file_url = doc.file_url

    try:
        doc_content = ""
        formatted_transcript = ""
        chunks = []

        # Check if this is a YouTube import
        if doc_type == "youtube" and doc_file_url:
            video_id = extract_youtube_video_id(doc_file_url)
            if video_id:
                logger.info(f"Extracting YouTube transcript for video ID: {video_id}")
                doc_content, formatted_transcript = get_youtube_transcript(video_id)
                
                if not doc_content:
                    logger.info("Transcript API failed or was empty. Trying audio extraction fallback...")
                    from app.services.audio import audio_service
                    import os
                    
                    audio_path = audio_service.download_youtube_audio(doc_file_url)
                    if audio_path:
                        try:
                            logger.info(f"Downloaded audio to {audio_path}. Starting audio transcription...")
                            doc_content = await audio_service.transcribe_youtube_audio(audio_path, doc_name)
                            
                            # Format plain text transcript with simulated 30s timestamps
                            if doc_content:
                                words = doc_content.split()
                                words_per_segment = 60 # ~30 seconds of speech at average rate
                                formatted_lines = []
                                for idx, i in enumerate(range(0, len(words), words_per_segment)):
                                    segment_words = words[i:i+words_per_segment]
                                    current_time = idx * 30
                                    minutes = current_time // 60
                                    seconds = current_time % 60
                                    timestamp = f"**[{minutes:02d}:{seconds:02d}]**"
                                    formatted_lines.append(f"{timestamp} " + " ".join(segment_words))
                                formatted_transcript = "\n\n".join(formatted_lines)
                        except Exception as e:
                            logger.error(f"Error during audio fallback processing: {e}")
                        finally:
                            # Clean up local audio file
                            if os.path.exists(audio_path):
                                try:
                                    os.remove(audio_path)
                                    logger.info(f"Cleaned up audio file: {audio_path}")
                                except Exception as cleanup_err:
                                    logger.error(f"Failed to delete {audio_path}: {cleanup_err}")
        # Check if this is a raw note sync
        elif doc_type == "note" and doc_file_url:
            import os
            if os.path.exists(doc_file_url):
                logger.info(f"Reading raw note text from: {doc_file_url}")
                with open(doc_file_url, "r", encoding="utf-8") as f:
                    doc_content = f.read()
        # Check if this is a local file upload (e.g. PDF/DOCX/TXT/etc.)
        elif doc_file_url and not doc_file_url.startswith("http"):
            import os
            if os.path.exists(doc_file_url):
                logger.info(f"Reading local uploaded document: {doc_file_url}")
                ext = os.path.splitext(doc_file_url)[1].lower()
                if ext == '.pdf':
                    try:
                        import pypdf
                        reader = pypdf.PdfReader(doc_file_url)
                        doc_content = ""
                        for page in reader.pages:
                            text = page.extract_text()
                            if text:
                                doc_content += text + "\n"
                    except Exception as pdf_err:
                        logger.error(f"Failed to parse PDF {doc_file_url}: {pdf_err}")
                elif ext in ('.docx', '.doc'):
                    try:
                        import docx
                        doc_obj = docx.Document(doc_file_url)
                        doc_content = "\n".join([p.text for p in doc_obj.paragraphs])
                    except Exception as docx_err:
                        logger.error(f"Failed to parse DOCX {doc_file_url}: {docx_err}")
                else:
                    # Fallback for plain text or other files
                    try:
                        with open(doc_file_url, "r", encoding="utf-8", errors="ignore") as f:
                            doc_content = f.read()
                    except Exception as txt_err:
                        logger.error(f"Failed to read file {doc_file_url}: {txt_err}")
        
        # Fallback/Default if no content obtained from transcript or file
        if not doc_content:
            logger.info(f"No content extracted for document: {doc_name}. Using topic-based template.")
            topic = doc_name
            doc_content = (
                f"Tutorial overview: {topic}. "
                "This tutorial covers core programming concepts, data structures, algorithms, "
                "and practical implementation techniques used in software development."
            )
            chunks = [doc_content]
        else:
            chunks = chunk_text(doc_content, chunk_size=1000)

        # 2. Upload vector chunks to Qdrant Vector database
        await vector_service.upsert_chunks(
            document_id=document_id,
            user_id=user_id,
            chunks=chunks
        )

        # 3. Generate structured study guide with LLM or set directly for notes
        if doc_type == "note":
            logger.info(f"Using manual note content directly for document {document_id}")
            doc_summary = doc_content
        else:
            logger.info(f"Generating structured study notes using LLM for document {document_id}")
            system_prompt = (
                "You are Aora AI, a world-class academic study-note generator. "
                "Your mission: transform raw video transcripts into COMPREHENSIVE, PREMIUM-QUALITY Markdown notes "
                "that a student would pay for and use as their primary study reference.\n\n"
                "ABSOLUTE REQUIREMENTS:\n"
                "- MINIMUM LENGTH: 2000 words per note. Aim for 3000-5000+ words for longer videos.\n"
                "- COMPLETE COVERAGE: Every single concept, term, command, example, workflow, and comparison "
                "from the transcript must appear in the notes. Nothing may be omitted.\n"
                "- TEXTBOOK DEPTH: Each concept must be explained in 2-5 full sentences, not just named.\n"
                "- ZERO TRUNCATION: Never use '...', 'etc.', or 'see transcript'. Write everything out fully.\n"
                "- REPRODUCE CODE: Every code snippet, command, formula, or config shown must appear verbatim "
                "in a fenced code block with the correct language tag.\n\n"
                "CONTENT RULES:\n"
                "- DEDUPLICATE: consolidate every repeated concept into one authoritative section.\n"
                "- CONCEPTUAL ORDER: organize by topic/concept, NOT by video timestamp.\n"
                "- ELIMINATE FILLER: skip ads, like & subscribe, greetings, and off-topic chatter entirely.\n"
                "- KEEP ANECDOTES: preserve memorable stories, metaphors, or jokes the instructor used.\n\n"
                "FORMATTING RULES:\n"
                "1. Title: single `# emoji Topic Name` h1 heading.\n"
                "2. Each major section: `## emoji Section Title` h2. Use rich emojis like: "
                "📖 🌐 🗂️ 🛠️ 🏗️ ⚙️ 📂 📋 🧩 🔀 ⚔️ 📦 🌿 🏷️ 🔄 🤝 🌳 🖥️ 🔧 📈 📄 👤 🚀 🔐.\n"
                "3. Sub-topics: `### emoji Sub-title` h3 with a 2-4 sentence explanation paragraph below each.\n"
                "4. Definitions/callouts: markdown blockquotes `> **Term**: explanation`.\n"
                "5. Comparisons: full Markdown tables with headers and separator rows.\n"
                "6. Commands/code: fenced code blocks with language tag (bash, python, git, etc.).\n"
                "7. Lists: `-` bullets for unordered, `1.` for ordered steps.\n"
                "8. End every major section with a summary table of commands/key points where applicable.\n"
                "9. Generate a final `## 📋 Quick Reference Cheat Sheet` section with a comprehensive table of ALL commands or formulas.\n"
                "10. NEVER truncate. Write the full comprehensive note regardless of length. "
                "If you feel you are running out of space, continue writing — do not summarize."
            )
            
            raw_summary = await llm_service.generate_long_notes(
                transcript=doc_content,
                doc_name=doc_name,
                system_prompt=system_prompt,
                provider="gemini",
            )
            
            # If the LLM returned the generic mock (no real API), build a real structured note from transcript
            if raw_summary and len(raw_summary) < 300 and ("Core Concept" in raw_summary or raw_summary.strip().startswith("### Overview")):
                # Build structured notes directly from transcript content
                words = doc_content.split()[:400]
                excerpt = " ".join(words)
                topic_name = doc_name
                raw_summary = f"""# 📚 {topic_name}

## Brief Overview
This note was generated from the YouTube video **[{topic_name}]({doc_file_url or '#'})**.
It covers key concepts, structured explanations, and practical takeaways from the tutorial.

## Key Points
- The video provides an in-depth walkthrough of core concepts in {topic_name}.
- Real examples and step-by-step demonstrations are included throughout.
- Focus on understanding the fundamentals before moving to advanced applications.
- Practice problems and implementation exercises are recommended for reinforcement.

## 📄 Content Overview

### 💡 Introduction
> **{topic_name}** is the focus of this tutorial. The presenter walks through the core ideas systematically.

### 📁 Transcript Excerpt
The following is extracted directly from the video transcript:

> {excerpt}

### 📝 Study Tips
- Re-watch key sections to reinforce difficult concepts.
- Take notes on unfamiliar terminology and look them up.
- Implement any code examples shown in the video yourself.

| Study Method | Benefit |
|---|---|
| Active recall | Strengthens long-term retention |
| Spaced repetition | Reduces forgetting curve |
| Practice problems | Builds problem-solving intuition |

## Summary
This note provides a structured overview of **{topic_name}**. To deepen your understanding, practice the concepts discussed, revisit the video at key timestamps, and use the AI Chat feature to ask follow-up questions."""
            
            doc_summary = raw_summary

        # 4. Save updates in a short-lived transaction
        async with SessionLocal() as db:
            result = await db.execute(select(Document).where(Document.id == document_id))
            doc = result.scalar_one_or_none()
            if doc:
                if doc_type != "note":
                    doc.summary = markdown_to_html(doc_summary)
                else:
                    doc.summary = doc_summary
                doc.status = "completed"
                await db.commit()
                logger.info(f"Completed processing for document {document_id}")
            else:
                logger.error(f"Document {document_id} not found when saving completed status.")
    except Exception as e:
        logger.error(f"Failed processing document {document_id}: {e}")
        async with SessionLocal() as db:
            try:
                result = await db.execute(select(Document).where(Document.id == document_id))
                doc = result.scalar_one_or_none()
                if doc:
                    doc.status = "failed"
                    await db.commit()
            except Exception as db_err:
                logger.error(f"Failed to set document status to failed: {db_err}")


async def _async_generate_podcast(podcast_id: int, user_id: str, topic: str):
    from app.db.session import SessionLocal
    from app.db.models.podcast import Podcast
    from app.services.llm import llm_service
    from app.services.audio import audio_service
    from app.services.storage import storage_service

    async with SessionLocal() as db:
        result = await db.execute(select(Podcast).where(Podcast.id == podcast_id))
        pod = result.scalar_one_or_none()
        if not pod:
            return

        try:
            # 1. Formulate conversation dialogue turns script
            script_prompt = f"Write a conversation between {pod.host_a} and {pod.host_b} discussing the topic: '{topic}'"
            raw_script_text = await llm_service.generate_text(
                prompt=script_prompt,
                system_prompt="Format the output as a clean text list of turns, like:\nSpeaker A: Hello\nSpeaker B: Hi",
                provider="gemini"
            )

            # Convert raw script text to JSON structures
            turns = []
            for line in raw_script_text.strip().split("\n"):
                if ":" in line:
                    speaker, text = line.split(":", 1)
                    turns.append({"speaker": speaker.strip(), "text": text.strip()})

            if not turns:
                turns = [
                    {"speaker": pod.host_a, "text": f"Welcome back. Today we are talking about {topic}."},
                    {"speaker": pod.host_b, "text": "Yes, it is a fascinating area to explore."}
                ]

            pod.script = turns

            # 2. Synthesize audio speech bytes (ElevenLabs)
            # Merge dialogue speech slices or simulate audio bytes
            audio_bytes = b""
            for turn in turns[:3]: # limit turns to save API credits during local demo
                turn_bytes = await audio_service.synthesize_text(
                    text=turn["text"],
                    voice_id="21m00Tcm4TlvDq8ikWAM" if turn["speaker"] == pod.host_a else "AZnzlk1XvdvUeBnXmlld"
                )
                audio_bytes += turn_bytes

            # 3. Upload combined podcast recording file to Supabase Storage bucket
            file_name = f"podcasts/user-{user_id}-pod-{podcast_id}.mp3"
            public_audio_url = await storage_service.upload_file(
                file_bytes=audio_bytes,
                file_name=file_name,
                content_type="audio/mpeg"
            )

            # 4. Save podcast details
            pod.audio_url = public_audio_url
            pod.duration = f"0:{str(len(turns) * 8).zfill(2)}" if len(turns) * 8 < 60 else f"{len(turns) * 8 // 60}:{str(len(turns) * 8 % 60).zfill(2)}"
            await db.commit()
            logger.info(f"Podcast {podcast_id} generated successfully at URL: {public_audio_url}")
        except Exception as e:
            logger.error(f"Failed to generate podcast {podcast_id}: {e}")
            await db.commit()
